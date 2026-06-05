#!/usr/bin/env python3
"""
TMJ Tutoring — quiet-tutor finder.

A "quiet" tutor is one who has not run a session in three weeks (>= 21 days).
Export is current as of Friday 29 May 2026, so that is treated as "today".

Reads tutors.csv (roster, source of truth) and sessions.csv (recent sessions),
matches the two on tutor name, and writes a list a person can act on:
who has gone quiet, how long, and how to reach them.

Every quiet tutor is identified by tutor_id (the hard requirement). Sessions
whose tutor name cannot be confidently tied to a roster tutor are surfaced on a
separate sheet rather than dropped — a name with no id is itself worth a look.

Writes one Excel workbook next to the CSVs, with two tabs:
    quiet_tutors_report.xlsx
        "Quiet Tutors"  — the action list, longest-quiet first
        "Needs Review"  — session names that didn't resolve to a tutor_id

Requires openpyxl (pip install openpyxl). Run from the folder with the CSVs:
    python quiet_tutors.py
"""

import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

TODAY = date(2026, 5, 29)          # "current as of Friday 29 May 2026"
QUIET_DAYS = 21                    # three weeks
HERE = Path(__file__).resolve().parent


def norm(name: str) -> str:
    """Lowercase, collapse whitespace, drop punctuation used in initials/apostrophes."""
    return " ".join(name.lower().replace(".", " ").replace("'", "").split())


def load_tutors(path: Path) -> list[dict]:
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def load_sessions(path: Path) -> list[dict]:
    rows = []
    with path.open(newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            r["date"] = datetime.strptime(r["date"].strip(), "%Y-%m-%d").date()
            rows.append(r)
    return rows


def match_tutor(session_name: str, subject: str, tutors: list[dict]) -> tuple[dict | None, str]:
    """
    Resolve a session's tutor_name to a roster tutor.
    Returns (tutor_row | None, how) where `how` records the decision made.

    Strategy, in order of confidence:
      1. exact normalised-name match
      2. abbreviated name ("J. Smith", "Sarah L.") -> first-token + surname-initial,
         disambiguated by subject when more than one roster tutor fits.
    """
    sname = norm(session_name)

    # 1. exact match
    exact = [t for t in tutors if norm(t["name"]) == sname]
    if len(exact) == 1:
        return exact[0], "exact"

    # 2. abbreviated form: a leading word + an initial somewhere (e.g. "j smith", "sarah l")
    tokens = sname.split()
    cands = []
    for t in tutors:
        tparts = norm(t["name"]).split()
        if len(tokens) != len(tparts):
            continue
        ok = True
        for st, tt in zip(tokens, tparts):
            # each session token must either equal the roster token or be its initial
            if st == tt or (len(st) == 1 and tt.startswith(st)):
                continue
            ok = False
            break
        if ok:
            cands.append(t)

    if len(cands) == 1:
        return cands[0], "abbreviated"
    if len(cands) > 1:
        by_subject = [t for t in cands if t["subject"].lower() == subject.strip().lower()]
        if len(by_subject) == 1:
            return by_subject[0], "abbreviated+subject"
        return None, f"ambiguous ({len(cands)} candidates)"

    return None, "no roster match"


def main() -> int:
    tutors = load_tutors(HERE / "tutors.csv")
    sessions = load_sessions(HERE / "sessions.csv")

    last_seen: dict[str, date] = {}      # tutor_id -> most recent session date
    session_count: dict[str, int] = {}   # tutor_id -> sessions matched
    unmatched: dict[str, dict] = {}      # raw session name -> info, for human review

    for s in sessions:
        tutor, how = match_tutor(s["tutor_name"], s["subject"], tutors)
        if tutor is None:
            u = unmatched.setdefault(
                s["tutor_name"],
                {"reason": how, "count": 0, "last": s["date"], "subjects": set()},
            )
            u["count"] += 1
            u["last"] = max(u["last"], s["date"])
            u["subjects"].add(s["subject"])
            continue
        tid = tutor["tutor_id"]
        session_count[tid] = session_count.get(tid, 0) + 1
        if tid not in last_seen or s["date"] > last_seen[tid]:
            last_seen[tid] = s["date"]

    # Build the quiet list. A roster tutor with no matched session at all is
    # quiet by definition (and noteworthy), so include them too.
    quiet = []
    for t in tutors:
        tid = t["tutor_id"]
        last = last_seen.get(tid)
        days = (TODAY - last).days if last else None
        if last is None or days >= QUIET_DAYS:
            quiet.append(
                {
                    "tutor_id": tid,
                    "name": t["name"],
                    "subject": t["subject"],
                    "phone": t["phone"],
                    "email": t["email"],
                    "last_session": last.isoformat() if last else "NEVER",
                    "days_quiet": days if days is not None else "n/a",
                    "sessions_seen": session_count.get(tid, 0),
                }
            )

    quiet.sort(key=lambda r: (r["days_quiet"] == "n/a", -(r["days_quiet"] if isinstance(r["days_quiet"], int) else 0)))

    # --- Output: one workbook, two tabs the office can open in Excel ---
    out = HERE / "quiet_tutors_report.xlsx"
    wb = Workbook()

    cols = ["tutor_id", "name", "subject", "last_session", "days_quiet", "phone", "email", "sessions_seen"]
    ws1 = wb.active
    ws1.title = "Quiet Tutors"
    ws1.append([c.replace("_", " ").title() for c in cols])
    for r in quiet:
        ws1.append([r[c] for c in cols])

    rcols = ["session_name", "reason", "sessions", "last_session", "subjects"]
    ws2 = wb.create_sheet("Needs Review")
    ws2.append([c.replace("_", " ").title() for c in rcols])
    for name, u in sorted(unmatched.items()):
        ws2.append([name, u["reason"], u["count"], u["last"].isoformat(),
                    ", ".join(sorted(u["subjects"]))])

    # light polish: bold headers, frozen top row, auto-ish column widths
    for ws in (ws1, ws2):
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for i, _ in enumerate(ws[1], start=1):
            letter = get_column_letter(i)
            width = max((len(str(c.value)) for c in ws[letter] if c.value is not None), default=10)
            ws.column_dimensions[letter].width = min(width + 2, 40)

    wb.save(out)

    # --- Also print a readable summary for the person doing the check-ins ---
    print(f"\nTMJ quiet-tutor check  -  as of {TODAY.isoformat()}  (quiet = no session in {QUIET_DAYS}+ days)\n")
    print(f"{len(quiet)} of {len(tutors)} tutors have gone quiet:\n")
    for r in quiet:
        print(f"  [{r['tutor_id']}] {r['name']:<16} {r['subject']:<10} "
              f"last seen {r['last_session']}  ({r['days_quiet']} days)  "
              f"{r['phone']}  {r['email']}")

    if unmatched:
        print(f"\nNeeds a human - {len(unmatched)} session name(s) not tied to a roster tutor_id:")
        for name, u in sorted(unmatched.items()):
            subs = ", ".join(sorted(u["subjects"]))
            print(f"  '{name}'  - {u['count']} session(s), last {u['last'].isoformat()}, "
                  f"subject {subs}  [{u['reason']}]")

    print(f"\nWritten to {out.name}  (tabs: 'Quiet Tutors', 'Needs Review')\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
